"""
Page numbering for DOCX documents
"""

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt
import logging

logger = logging.getLogger(__name__)

# Alignment mapping for CLI argument
ALIGNMENT_MAP = {
    'center': WD_PARAGRAPH_ALIGNMENT.CENTER,
    'left': WD_PARAGRAPH_ALIGNMENT.LEFT,
    'right': WD_PARAGRAPH_ALIGNMENT.RIGHT,
}

# Page number format options
PAGE_NUMBER_FORMATS = {
    'simple': '{page}',
    'page_x': 'Page {page}',
    'page_x_of_y': 'Page {page} of {total}',
    'x_of_y': '{page} of {total}',
    'dash': '- {page} -',
    'dash_total': '- {page} of {total} -',
}


def create_page_field() -> tuple:
    """
    Create a PAGE field element for current page number

    Returns:
        tuple: The PAGE field elements
    """
    fldChar_begin = OxmlElement('w:fldChar')
    fldChar_begin.set(qn('w:fldCharType'), 'begin')

    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = "PAGE"

    fldChar_separate = OxmlElement('w:fldChar')
    fldChar_separate.set(qn('w:fldCharType'), 'separate')

    fldChar_end = OxmlElement('w:fldChar')
    fldChar_end.set(qn('w:fldCharType'), 'end')

    return fldChar_begin, instrText, fldChar_separate, fldChar_end


def create_numpages_field() -> tuple:
    """
    Create a NUMPAGES field element for total page count

    Returns:
        tuple: The NUMPAGES field elements
    """
    fldChar_begin = OxmlElement('w:fldChar')
    fldChar_begin.set(qn('w:fldCharType'), 'begin')

    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = "NUMPAGES"

    fldChar_separate = OxmlElement('w:fldChar')
    fldChar_separate.set(qn('w:fldCharType'), 'separate')

    fldChar_end = OxmlElement('w:fldChar')
    fldChar_end.set(qn('w:fldCharType'), 'end')

    return fldChar_begin, instrText, fldChar_separate, fldChar_end


def add_field_to_run(run, field_elements: tuple) -> None:
    """
    Add field elements to a run

    Args:
        run: Document run
        field_elements: Tuple of field elements
    """
    for element in field_elements:
        run._r.append(element)


def add_page_numbers(
        docx_path: str,
        skip_first_page: bool = True,
        alignment: str | WD_PARAGRAPH_ALIGNMENT = 'right',
        font_name: str = 'Arial',
        font_size: int = 10,
        page_number_format: str = 'simple',
        custom_format: str | None = None,
) -> None:
    """
    Add page numbers to a DOCX document

    Args:
        docx_path: Path to the DOCX file
        skip_first_page: Whether to skip numbering on the first page (title page)
        alignment: Alignment of page numbers ('center', 'left', 'right' or WD_PARAGRAPH_ALIGNMENT)
        font_name: Font name for page numbers
        font_size: Font size for page numbers
        page_number_format: Predefined format name ('simple', 'page_x', 'page_x_of_y', 'x_of_y', 'dash', 'dash_total')
        custom_format: Custom format string with {page} and {total} placeholders (overrides page_number_format)
    """
    try:
        doc = Document(docx_path)

        # Get the first section
        section = doc.sections[0]

        # Resolve alignment
        if isinstance(alignment, str):
            alignment = ALIGNMENT_MAP.get(alignment.lower(),
                                          WD_PARAGRAPH_ALIGNMENT.CENTER)

        # Resolve format string
        if custom_format:
            format_string = custom_format
        else:
            format_string = PAGE_NUMBER_FORMATS.get(page_number_format, '{page}')

        if skip_first_page:
            # Enable "Different First Page" option
            sectPr = section._sectPr

            # Remove existing titlePg if present
            existing_titlePg = sectPr.find(qn('w:titlePg'))
            if existing_titlePg is not None:
                sectPr.remove(existing_titlePg)

            titlePg = OxmlElement('w:titlePg')
            sectPr.append(titlePg)

        # Get the footer
        footer = section.footer
        footer.is_linked_to_previous = False

        # Clear existing footer content
        for paragraph in footer.paragraphs:
            paragraph.clear()

        # Get or create paragraph
        if footer.paragraphs:
            paragraph = footer.paragraphs[0]
        else:
            paragraph = footer.add_paragraph()

        paragraph.alignment = alignment

        # Parse format string and build footer content
        _build_page_number_content(paragraph, format_string, font_name,
                                   font_size)

        # Save the document
        doc.save(docx_path)
        logger.info(
            f"Page numbers added to: {docx_path} (format: {format_string})")

    except Exception as e:
        logger.error(f"Failed to add page numbers to {docx_path}: {e}")
        raise


def _build_page_number_content(
        paragraph,
        format_string: str,
        font_name: str,
        font_size: int,
) -> None:
    """
    Build page number content in a paragraph based on format string

    Args:
        paragraph: Document paragraph
        format_string: Format string with {page} and/or {total} placeholders
        font_name: Font name
        font_size: Font size
    """
    # Track position in format string
    current_pos = 0
    format_len = len(format_string)

    while current_pos < format_len:
        # Find next placeholder
        page_pos = format_string.find('{page}', current_pos)
        total_pos = format_string.find('{total}', current_pos)

        # Determine which placeholder comes first
        if page_pos == -1 and total_pos == -1:
            # No more placeholders, add remaining text
            remaining_text = format_string[current_pos:]
            if remaining_text:
                run = paragraph.add_run(remaining_text)
                run.font.name = font_name
                run.font.size = Pt(font_size)
            break

        # Find the nearest placeholder
        if page_pos == -1:
            next_pos = total_pos
            placeholder = '{total}'
        elif total_pos == -1:
            next_pos = page_pos
            placeholder = '{page}'
        else:
            if page_pos < total_pos:
                next_pos = page_pos
                placeholder = '{page}'
            else:
                next_pos = total_pos
                placeholder = '{total}'

        # Add text before placeholder
        if next_pos > current_pos:
            text_before = format_string[current_pos:next_pos]
            run = paragraph.add_run(text_before)
            run.font.name = font_name
            run.font.size = Pt(font_size)

        # Add the field
        run = paragraph.add_run()
        run.font.name = font_name
        run.font.size = Pt(font_size)

        if placeholder == '{page}':
            add_field_to_run(run, create_page_field())
        else:  # {total}
            add_field_to_run(run, create_numpages_field())

        # Move past the placeholder
        current_pos = next_pos + len(placeholder)


def get_available_formats() -> dict:
    """
    Get available page number format options

    Returns:
        dict: Dictionary of format names and their patterns
    """
    return PAGE_NUMBER_FORMATS.copy()


def get_available_alignments() -> list:
    """
    Get available alignment options

    Returns:
        list: List of alignment option names
    """
    return list(ALIGNMENT_MAP.keys())
