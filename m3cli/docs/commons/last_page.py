"""
Last page generator for DOCX documents - includes Changelog and Version info
"""

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT, WD_BREAK
import logging
import subprocess
import sys

logger = logging.getLogger(__name__)

# Version History table settings
VERSION_TABLE_HEADERS = ['Version', 'Date', 'Summary']
VERSION_TABLE_EMPTY_ROWS = 2


def get_cli_version(
        cli_entry_point: str,
        file_path: str | None = None,
        module_path: str | None = None,
) -> str | None:
    """
    Get CLI version by running 'cli_entry_point --version' command

    Args:
        cli_entry_point: Name of the CLI entry point (e.g., 'm3')
        file_path: Path to the CLI file (optional)
        module_path: Python module path (optional)

    Returns:
        Version string or None if unable to determine
    """
    try:
        # Try direct command first (if CLI is installed)
        try:
            result = subprocess.run(
                [cli_entry_point, '--version'],
                capture_output=True,
                text=True,
                timeout=5,
            )
            if result.returncode == 0:
                output = result.stdout.strip()
                logger.info(f"CLI version output: {output}")
                return output
        except (FileNotFoundError, subprocess.TimeoutExpired):
            pass

        # Try running as Python module
        if file_path:
            try:
                result = subprocess.run(
                    [sys.executable, file_path, '--version'],
                    capture_output=True,
                    text=True,
                    timeout=5,
                )
                if result.returncode == 0:
                    output = result.stdout.strip()
                    logger.info(f"CLI version output: {output}")
                    return output
            except subprocess.TimeoutExpired:
                pass

        # Try as Python module import
        if module_path:
            try:
                result = subprocess.run(
                    [sys.executable, '-m', module_path, '--version'],
                    capture_output=True,
                    text=True,
                    timeout=5,
                )
                if result.returncode == 0:
                    output = result.stdout.strip()
                    logger.info(f"CLI version output: {output}")
                    return output
            except subprocess.TimeoutExpired:
                pass

        logger.warning(
            f"Unable to determine CLI version for '{cli_entry_point}'")
        return None

    except Exception as e:
        logger.error(f"Error getting CLI version: {e}")
        return None


def add_hyperlink(paragraph, text: str, url: str) -> None:
    """
    Add a hyperlink to a paragraph

    Args:
        paragraph: Document paragraph
        text: Display text for the link
        url: URL target
    """
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True
    )

    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)

    new_run = OxmlElement('w:r')

    text_elem = OxmlElement('w:t')
    text_elem.text = text
    new_run.append(text_elem)

    rPr = OxmlElement('w:rPr')

    color = OxmlElement('w:color')
    color.set(qn('w:val'), '0563C1')
    rPr.append(color)

    u = OxmlElement('w:u')
    u.set(qn('w:val'), 'single')
    rPr.append(u)

    new_run.insert(0, rPr)
    hyperlink.append(new_run)

    paragraph._p.append(hyperlink)


def set_cell_border(cell, border_color: str = 'CCCCCC') -> None:
    """Set borders for a table cell"""
    tc_pr = cell._element.get_or_add_tcPr()
    tc_borders = OxmlElement('w:tcBorders')

    for side in ['top', 'left', 'bottom', 'right']:
        border = OxmlElement(f'w:{side}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), border_color)
        tc_borders.append(border)

    tc_pr.append(tc_borders)


def set_cell_background(cell, color: str) -> None:
    """Set background color for a table cell"""
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), color)
    cell._element.get_or_add_tcPr().append(shd)


def ensure_document_section_settings(doc: Document) -> None:
    """
    Ensure document has valid section settings (page width, margins).
    This is needed for add_table() to work properly.
    """
    for section in doc.sections:
        # Set default page size if not set (Letter size: 8.5 x 11 inches)
        if section.page_width is None:
            section.page_width = Inches(8.5)
        if section.page_height is None:
            section.page_height = Inches(11)

        # Set default margins if not set
        if section.left_margin is None:
            section.left_margin = Inches(0.4)
        if section.right_margin is None:
            section.right_margin = Inches(0.4)
        if section.top_margin is None:
            section.top_margin = Inches(0.4)
        if section.bottom_margin is None:
            section.bottom_margin = Inches(0.4)


def add_version_history_table(doc: Document) -> None:
    """
    Add a version history table to the document
    """
    logger.debug("Creating version history table...")

    headers = VERSION_TABLE_HEADERS
    empty_rows = VERSION_TABLE_EMPTY_ROWS

    # Ensure document has valid section settings before adding table
    ensure_document_section_settings(doc)

    # Create table with header row + empty rows
    total_rows = 1 + empty_rows
    table = doc.add_table(rows=total_rows, cols=len(headers))

    logger.debug(f"Table created with {total_rows} rows and {len(headers)} columns")

    # Style header row
    header_row = table.rows[0]
    for i, header_text in enumerate(headers):
        cell = header_row.cells[i]

        # Clear existing content and add header text
        cell.text = ''
        para = cell.paragraphs[0]
        run = para.add_run(header_text)
        run.font.name = 'Arial'
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
        para.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT

        # Style header cell
        set_cell_border(cell)
        set_cell_background(cell, 'E0E0E0')

    logger.debug("Header row styled")

    # Style empty data rows
    for row_idx in range(1, total_rows):
        row = table.rows[row_idx]
        for cell in row.cells:
            set_cell_border(cell)
            # Ensure there's a paragraph
            if not cell.paragraphs:
                cell.add_paragraph()
            para = cell.paragraphs[0]
            para.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
            # Add non-breaking space to ensure row has height
            run = para.add_run('\u00A0')
            run.font.name = 'Arial'
            run.font.size = Pt(10)

    logger.debug("Data rows styled")
    logger.info("Version history table added successfully")


def add_last_page_to_docx(
        docx_path: str,
        cli_version: str | None = None,
        changelog_url: str | None = None,
        changelog_path: str | None = None,
) -> None:
    """
    Add a last page with Changelog and Version information to a DOCX document

    Args:
        docx_path: Path to the DOCX file
        cli_version: CLI version string (e.g., "m3, version 4.158.0")
        changelog_url: URL to the changelog (e.g., Git repository URL)
        changelog_path: Local path to changelog file
    """
    try:
        logger.debug(f"Opening document: {docx_path}")
        doc = Document(docx_path)

        # Ensure document has valid section settings
        ensure_document_section_settings(doc)

        # Add page break to create a new page
        logger.debug("Adding page break...")
        last_para = doc.paragraphs[-1] if doc.paragraphs else doc.add_paragraph()
        run = last_para.add_run()
        run.add_break(WD_BREAK.PAGE)

        # Add "Version History" heading
        logger.debug("Adding Version History heading...")
        version_history_heading = doc.add_paragraph()
        run = version_history_heading.add_run("Version History")
        run.font.name = 'Arial'
        run.font.size = Pt(20)
        run.font.bold = True
        run.font.color.rgb = RGBColor(70, 69, 71)
        version_history_heading.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT

        # Add spacing before table
        doc.add_paragraph()

        # Add version history table
        logger.debug("Adding version history table...")
        add_version_history_table(doc)

        # Add spacing after table
        doc.add_paragraph()

        # Add changelog link/path if provided
        if changelog_url or changelog_path:
            logger.debug("Adding changelog link...")
            link_para = doc.add_paragraph()
            link_para.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT

            if changelog_url:
                run = link_para.add_run("View Changelog: ")
                run.font.name = 'Arial'
                run.font.size = Pt(10)

                add_hyperlink(link_para, changelog_url, changelog_url)
            elif changelog_path:
                run = link_para.add_run("Changelog location: ")
                run.font.name = 'Arial'
                run.font.size = Pt(10)

                path_run = link_para.add_run(changelog_path)
                path_run.font.name = 'Arial'
                path_run.font.size = Pt(10)
                path_run.font.color.rgb = RGBColor(0, 0, 0)

        # Add single empty line spacing
        doc.add_paragraph()

        # Add version information if provided
        if cli_version:
            logger.debug("Adding version info...")
            version_para = doc.add_paragraph()

            run = version_para.add_run("The Documentation generated for: ")
            run.font.name = 'Arial'
            run.font.size = Pt(10)

            version_run = version_para.add_run(cli_version)
            version_run.font.name = 'Arial'
            version_run.font.size = Pt(10)
            version_run.font.bold = True
            version_run.font.color.rgb = RGBColor(0, 0, 0)

            version_para.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT

        logger.debug("Saving document...")
        doc.save(docx_path)
        logger.info(f"Last page successfully added to: {docx_path}")

    except Exception as e:
        import traceback
        logger.error(f"Failed to add last page to {docx_path}: {e}")
        logger.error(f"Traceback: {traceback.format_exc()}")
        raise


def add_cli_last_page(
        docx_path: str,
        cli_name: str,
        cli_version: str | None = None,
        changelog_url: str | None = None,
        changelog_path: str | None = None,
        auto_detect_version: bool = True,
        cli_file_path: str | None = None,
        cli_module_path: str | None = None,
) -> None:
    """
    Add a last page for CLI documentation with auto-detection of version

    Args:
        docx_path: Path to the DOCX file
        cli_name: Name of the CLI tool (e.g., 'm3')
        cli_version: CLI version string (optional, will auto-detect if not provided)
        changelog_url: URL to the changelog
        changelog_path: Local path to changelog file
        auto_detect_version: Whether to auto-detect version if not provided
        cli_file_path: Path to CLI file for version detection
        cli_module_path: Module path for version detection
    """
    if not cli_version and auto_detect_version:
        logger.info(f"Attempting to auto-detect version for '{cli_name}'...")
        cli_version = get_cli_version(
            cli_entry_point=cli_name,
            file_path=cli_file_path,
            module_path=cli_module_path,
        )

        if cli_version:
            logger.info(f"Auto-detected version: {cli_version}")
        else:
            logger.warning(f"Could not auto-detect version for '{cli_name}'")

    add_last_page_to_docx(
        docx_path=docx_path,
        cli_version=cli_version,
        changelog_url=changelog_url,
        changelog_path=changelog_path,
    )
